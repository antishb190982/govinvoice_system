"""Bounded statement parsing; identifiers are never fuzzy-matched or numerically rounded."""
import csv
import io
import re
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.exceptions import ValidationError

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10000
ALIASES = {
    "account_number": {"account", "account number", "account no", "service account", "customer account", "customer number"},
    "invoice_number": {"invoice", "invoice number", "invoice no", "bill number", "bill no", "document number"},
    "amount": {"amount", "amount total", "amount due", "pending amount", "outstanding amount", "invoice amount", "total"},
    "currency": {"currency", "ccy"},
}


def norm(value):
    return str(value or "").strip().casefold()


def text(value):
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if value == int(value):
            return str(int(value))
    return str(value if value is not None else "").strip()


def parse_amount(value):
    raw = text(value)
    if not re.fullmatch(r"(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?", raw):
        raise ValueError("Amount must be a non-negative number, e.g. 1234.50; credits and account-level balances need investigation.")
    try:
        result = Decimal(raw.replace(",", ""))
        if result > Decimal("9999999999999.999") or result != result.quantize(Decimal(".001")):
            raise ValueError("Amount exceeds invoice precision (13 digits and 3 decimal places).")
        return result
    except InvalidOperation as exc:
        raise ValueError("Invalid amount.") from exc


def parse_statement(data, filename, currency):
    if len(data) > MAX_BYTES:
        raise ValidationError("Statement files must be at most 10 MB.")
    ext = Path(filename).suffix.lower()
    book = None
    try:
        if ext == ".csv":
            content = data.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            records = csv.reader(io.StringIO(content), dialect)
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValidationError("The expanded workbook exceeds 40 MB.")
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            records = book.active.iter_rows(values_only=True)
        elif ext == ".xls":
            import xlrd
            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            sheet = book.sheet_by_index(0)
            records = (sheet.row_values(i) for i in range(min(sheet.nrows, MAX_ROWS + 2)))
            if sheet.nrows > MAX_ROWS + 1:
                raise ValidationError("A statement may contain at most 10,000 invoice rows.")
        else:
            raise ValidationError("Upload a CSV, XLS or XLSX statement.")
        headers = next(records, [])
        headers = [re.sub(r"[_\W]+", " ", norm(h)).strip() for h in headers]
        mapping = {}
        for key, aliases in ALIASES.items():
            positions = [i for i, h in enumerate(headers) if h in aliases]
            if len(positions) > 1:
                raise ValidationError(f"More than one column maps to {key}; use the supplied template.")
            if positions:
                mapping[key] = positions[0]
        if not {"account_number", "invoice_number", "amount"}.issubset(mapping):
            raise ValidationError("Required columns: account_number, invoice_number, amount. Supply one row per invoice, not account totals.")
        rows = []
        for number, values in enumerate(records, 2):
            if number > MAX_ROWS + 1:
                raise ValidationError("A statement may contain at most 10,000 invoice rows.")
            if not any(v not in (None, "") for v in values):
                continue
            raw = {key: text(values[i]) if i < len(values) else "" for key, i in mapping.items()}
            errors = []
            for key in ["account_number", "invoice_number"]:
                if not raw[key] or len(raw[key]) > 100 or raw[key].startswith(("=", "+", "@")):
                    errors.append(f"{key} is missing, too long or a formula. Preserve identifiers as text, including leading zeros.")
            if raw.get("currency") and raw["currency"].upper() != currency:
                errors.append("Currency differs from the statement currency; no conversion is performed.")
            amount = None
            try:
                amount = parse_amount(raw["amount"])
            except ValueError as exc:
                errors.append(str(exc))
            rows.append({"line_number": number, "account_number": raw["account_number"][:100],
                         "invoice_number": raw["invoice_number"][:100], "amount": amount,
                         "raw": raw, "parse_error": " ".join(errors)})
        if not rows:
            raise ValidationError("The statement contains no invoice rows.")
        return rows
    except ValidationError:
        raise
    except Exception as exc:
        raise ValidationError("The statement could not be read. Use an unencrypted CSV, XLS or XLSX with the template columns.") from exc
    finally:
        if book:
            if ext == ".xlsx":
                book.close()
            else:
                book.release_resources()
