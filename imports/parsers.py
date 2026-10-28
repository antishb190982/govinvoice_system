"""Readers for provider spreadsheets and PDF invoices, producing normalised rows."""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

FIELDS = ["account_number", "service_number", "invoice_number", "issue_date", "due_date", "period_start",
          "period_end", "amount_net", "amount_tax", "amount_total", "arrears", "consumption", "previous_reading",
          "current_reading", "currency", "provider"]
# Read from files but not stored as such: a billing month like "September 2026"
# or "01/09/2026 - 30/09/2026" that fills period_start / period_end.
EXTRA_FIELDS = ["period"]
DATE_FIELDS = {"issue_date", "due_date", "period_start", "period_end"}
DECIMAL_FIELDS = {"amount_net", "amount_tax", "amount_total", "arrears", "consumption", "previous_reading",
                  "current_reading"}
REQUIRED = {"account_number", "invoice_number", "issue_date", "due_date", "amount_total"}

ALIASES = {
    "account_number": ["account no id", "account", "account no", "account number", "acct", "contract", "contract no",
                       "contract acct", "contract account", "billing account", "bill acc", "bill account",
                       "customer account", "subscriber no", "premise no", "meter account"],
    "service_number": ["telephone number", "telephone no", "line number", "phone number", "service number",
                       "meter no", "meter number", "msisdn"],
    "invoice_number": ["bill number", "invoice", "invoice no", "invoice number", "bill no", "document no",
                       "inv no", "reference", "bill ref", "serial no"],
    "issue_date": ["issue date", "invoice date", "bill date", "bill issued on", "date", "billing date", "document date"],
    "due_date": ["due date", "payment due", "pay by", "due", "last payment date", "last date for payment"],
    "period_start": ["period start", "from", "from date", "period from", "start date"],
    "period_end": ["period end", "to", "to date", "period to", "end date"],
    "amount_net": ["net", "net amount", "amount excl vat", "subtotal", "amount before tax"],
    "amount_tax": ["current vat bill", "vat", "tax", "vat amount", "tax amount", "value added tax"],
    "amount_total": ["current amount bill incl vat", "current amount bill", "current bill", "current amount payable",
                     "total", "amount", "total amount", "amount due", "current amount due", "total due",
                     "net payable", "invoice amount", "gross", "amount incl vat", "total payable"],
    "arrears": ["arrears", "outstanding", "outstanding amount", "outstanding amt", "balance brought forward",
                "previous balance", "balance from previous charges"],
    "period": ["period", "billing period", "billing month", "month", "consumption period"],
    "consumption": ["consumption", "units", "usage", "kwh", "quantity", "volume", "m3"],
    "previous_reading": ["previous reading", "prev reading", "opening reading"],
    "current_reading": ["current reading", "present reading", "closing reading"],
    "currency": ["currency", "ccy"],
    "provider": ["provider", "supplier", "company"],
}

DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d.%m.%y", "%d-%m-%y", "%d %b %Y",
                "%d-%b-%Y", "%d %B %Y", "%b %d, %Y", "%Y/%m/%d"]
MONTH_RE = re.compile(r"^\s*([A-Za-zéû]+)\.?\s+(\d{4})\s*$")
FRENCH_MONTHS = {"janvier": 1, "février": 2, "fevrier": 2, "mars": 3, "avril": 4, "mai": 5, "juin": 6, "juillet": 7,
                 "août": 8, "aout": 8, "septembre": 9, "octobre": 10, "novembre": 11, "décembre": 12, "decembre": 12}


def parse_period(value):
    """'September 2026' → (1 Sep, 30 Sep); '01/09/2026 - 30/09/2026' → both dates; else (None, None)."""
    import calendar

    if value in (None, ""):
        return None, None
    text = str(value).strip()
    m = MONTH_RE.match(text)
    if m:
        name, year = m.group(1).lower(), int(m.group(2))
        month = FRENCH_MONTHS.get(name)
        if month is None:
            for fmt in ("%B", "%b"):
                try:
                    month = datetime.strptime(name.capitalize(), fmt).month
                    break
                except ValueError:
                    continue
        if month:
            return date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1])
    parts = re.split(r"\s+(?:-|–|to|au|à)\s+|\s*-\s*(?=\d{1,2}[/.])", text)
    if len(parts) == 2:
        try:
            return parse_date(parts[0]), parse_date(parts[1])
        except ValueError:
            pass
    return None, None


def _norm(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def match_headers(headers, column_map=None):
    """Map our field names to column indexes."""
    mapping = {}
    normalised = [_norm(h) for h in headers]
    if column_map:
        for field, header in column_map.items():
            if field in FIELDS + EXTRA_FIELDS and _norm(header) in normalised:
                mapping[field] = normalised.index(_norm(header))
    for field in FIELDS + EXTRA_FIELDS:
        if field in mapping:
            continue
        for alias in ALIASES[field]:
            if alias in normalised and normalised.index(alias) not in mapping.values():
                mapping[field] = normalised.index(alias)
                break
    return mapping


def parse_date(value, fmt=None):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        pass
    for f in ([fmt] if fmt else []) + DATE_FORMATS:
        try:
            return datetime.strptime(text, f).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date “{text}”")


def parse_decimal(value, decimal_comma=False):
    if value in (None, ""):
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.001"))
    text = re.sub(r"[^\d,.\-()]", "", str(value))
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    if decimal_comma:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        number = Decimal(text).quantize(Decimal("0.001"))
    except InvalidOperation:
        raise ValueError(f"unrecognised number “{value}”")
    return -number if negative else number


def normalise(raw, date_format=None, decimal_comma=False):
    """Convert raw strings to typed values; returns (data, errors)."""
    data, errors = {}, []
    for field in FIELDS:
        value = raw.get(field)
        try:
            if field in DATE_FIELDS:
                parsed = parse_date(value, date_format)
                data[field] = parsed.isoformat() if parsed else None
            elif field in DECIMAL_FIELDS:
                parsed = parse_decimal(value, decimal_comma)
                data[field] = str(parsed) if parsed is not None else None
            else:
                data[field] = str(value).strip() if value not in (None, "") else None
                if isinstance(value, float) and value.is_integer():
                    data[field] = str(int(value))
        except ValueError as exc:
            errors.append(f"{field.replace('_', ' ')}: {exc}")
            data[field] = None
    if raw.get("period") and not (data.get("period_start") and data.get("period_end")):
        start, end = parse_period(raw["period"])
        if start:
            data["period_start"], data["period_end"] = data.get("period_start") or start.isoformat(), \
                data.get("period_end") or end.isoformat()
        else:
            errors.append(f"period: unrecognised “{raw['period']}”")
    for field in REQUIRED:
        if not data.get(field) and not any(e.startswith(field.replace("_", " ")) for e in errors):
            errors.append(f"{field.replace('_', ' ')} is missing")
    if data.get("account_number"):
        data["account_number"] = data["account_number"].strip()
    if data.get("amount_total") and Decimal(data["amount_total"]) < 0:
        errors.append("amount total is negative (credit notes are not imported)")
    return data, errors


def read_table(content, kind, sheet_name=""):
    """Return rows (lists of cell values) from CSV/XLSX/XLS bytes."""
    if kind == "csv":
        text = content.decode("utf-8-sig", errors="replace")
        # Count candidate delimiters in the header line; the sniffer is unreliable on
        # exports with trailing separators and padded headers (e.g. Mauritius Telecom).
        first = next((ln for ln in text.splitlines() if ln.strip()), "")
        delimiter = max(",;\t|", key=first.count) if first else ","
        if not first.count(delimiter):
            delimiter = ","
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    if kind == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if kind == "xls":
        import xlrd

        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_name(sheet_name) if sheet_name and sheet_name in book.sheet_names() else book.sheet_by_index(0)
        rows = []
        for r in range(sheet.nrows):
            row = []
            for c in range(sheet.ncols):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
                else:
                    row.append(cell.value)
            rows.append(row)
        return rows
    raise ValueError(f"unsupported table kind {kind}")


def spreadsheet_rows(content, kind, template=None):
    rows = read_table(content, kind, template.sheet_name if template else "")
    header_index = (template.header_row - 1) if template else 0
    if not template:
        # Find the first row that looks like a header.
        for i, row in enumerate(rows[:20]):
            if len(match_headers(row)) >= 3:
                header_index = i
                break
    if header_index >= len(rows):
        raise ValueError("The file has no header row.")
    headers = rows[header_index]
    mapping = match_headers(headers, template.column_map if template else None)
    missing = [f for f in REQUIRED if f not in mapping]
    if missing:
        raise ValueError("Could not find columns for: " + ", ".join(m.replace("_", " ") for m in sorted(missing)) +
                         ". Add an import template for this provider.")
    out = []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(c not in (None, "") for c in row):
            continue
        raw = {f: (row[i] if i < len(row) else None) for f, i in mapping.items()}
        raw = {k: (v.strip() if isinstance(v, str) else v) for k, v in raw.items()}
        if _norm(raw.get("account_number")) in {"total", "grand total"}:
            continue
        out.append((offset, {k: (v.isoformat() if isinstance(v, (date, datetime)) else v) for k, v in raw.items()}))
    return out


PDF_DEFAULT_PATTERNS = {
    "invoice_number": r"(?:invoice|bill)\s*(?:no|number|#)\.?\s*[:#]?\s*([A-Z0-9][A-Z0-9\-/]{2,})",
    "account_number": r"(?:account|contract|customer)\s*(?:no|number|#)\.?\s*[:#]?\s*([A-Z0-9][A-Z0-9\-/]{2,})",
    "issue_date": r"(?:invoice|bill|issue)\s*date\s*[:\-]?\s*([0-9]{1,4}[\-/. ][0-9A-Za-z]{1,9}[\-/. ,]+[0-9]{2,4})",
    "due_date": r"(?:due\s*date|pay\s*by|payment\s*due)\s*[:\-]?\s*([0-9]{1,4}[\-/. ][0-9A-Za-z]{1,9}[\-/. ,]+[0-9]{2,4})",
    "period_start": r"period\s*(?:from)?\s*[:\-]?\s*([0-9]{1,4}[\-/.][0-9]{1,2}[\-/.][0-9]{2,4})\s*(?:to|-|–)",
    "period_end": r"period\s*(?:from)?\s*[:\-]?\s*[0-9]{1,4}[\-/.][0-9]{1,2}[\-/.][0-9]{2,4}\s*(?:to|-|–)\s*([0-9]{1,4}[\-/.][0-9]{1,2}[\-/.][0-9]{2,4})",
    "amount_tax": r"(?:vat|tax)\s*(?:amount)?\s*(?:\(\s*\d+\s*%\s*\))?\s*[:\-]?\s*(?:[A-Z]{3}\s*)?([0-9][0-9,]*\.?[0-9]*)",
    "amount_total": r"(?:total\s*(?:amount)?\s*(?:due|payable)|amount\s*due|net\s*payable)\s*[:\-]?\s*(?:[A-Z]{3}\s*)?([0-9][0-9,]*\.?[0-9]*)",
    "consumption": r"(?:consumption|units\s*consumed|usage)\s*[:\-]?\s*([0-9][0-9,]*\.?[0-9]*)",
    "previous_reading": r"previous\s*reading\s*[:\-]?\s*([0-9][0-9,]*\.?[0-9]*)",
    "current_reading": r"(?:current|present)\s*reading\s*[:\-]?\s*([0-9][0-9,]*\.?[0-9]*)",
    "currency": r"\b(OMR|AED|SAR|QAR|KWD|BHD|USD|EUR|GBP)\b",
}


def pdf_text(content):
    import pdfplumber

    with pdfplumber.open(io.BytesIO(content)) as pdf:
        if len(pdf.pages) > 50:
            raise ValueError("PDF has more than 50 pages; split it first.")
        return "\n".join((page.extract_text() or "") for page in pdf.pages)


def pdf_row(content, template=None):
    text = pdf_text(content)
    if len(text.strip()) < 20:
        raise ValueError("The PDF has no text layer (scanned image). Enter it manually or enable OCR.")
    patterns = dict(PDF_DEFAULT_PATTERNS)
    if template and template.pdf_patterns:
        patterns.update(template.pdf_patterns)
    raw = {}
    for field, pattern in patterns.items():
        match = re.search(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
        if match:
            raw[field] = match.group(1).strip()
    return raw, text
